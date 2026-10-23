"""Tkinter front end with one worker and main-thread-only widget updates."""

from __future__ import annotations

import argparse
import queue
import threading
import traceback
from collections.abc import Callable
from tkinter import END, Listbox, StringVar, Text, Tk, Toplevel, filedialog, messagebox, ttk

from . import __version__
from .checks import check_environment
from .cli import build_arg_parser, resolve_paths
from .converter import convert


class _GuiParser(argparse.ArgumentParser):
    def error(self, message: str) -> None:
        raise ValueError(message)


class Application:
    def __init__(self, root: Tk, args: argparse.Namespace | None = None) -> None:
        self.root = root
        self.busy = False
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.schema_files = list(args.schema or []) if args else []
        self.sheets = {
            side: list(getattr(args, side + "_sheet", []) or []) for side in ("input", "output")
        }
        self.fields = {
            key: StringVar(master=root, value=getattr(args, key, None) or "" if args else "")
            for key in (
                "excel",
                "dbc_folder",
                "blf",
                "output",
                "config",
                "matlab",
                "element_column",
                "can_column",
                "interface_column",
                "dtype_column",
                "header_row",
                "output_format",
            )
        }
        if not self.fields["output_format"].get():
            self.fields["output_format"].set("dataset" if args and args.matlab else "mat")
        self.schema_label = StringVar(master=root, value="; ".join(self.schema_files))
        self.sheet_label = StringVar(
            master=root,
            value=" | ".join(
                f"{side}: {', '.join(names)}" for side, names in self.sheets.items() if names
            )
            or "Auto-detect sheets",
        )
        self.status = StringVar(master=root, value="Ready")
        self.debug = bool(args and args.debug)
        root.title(f"BLF to MAT {__version__}")
        root.minsize(900, 720)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(0, weight=1)
        frame = ttk.Frame(root, padding=12)
        frame.grid(sticky="nsew")
        frame.columnconfigure(1, weight=1)
        self.controls: list[ttk.Widget] = []
        rows = (
            ("Excel workbook (optional)", "excel", "excel"),
            ("DBC file / folder (optional)", "dbc_folder", "dbc"),
            ("BLF file / folder", "blf", "blf"),
            ("TXT schemas (optional)", "schema", "schema"),
            ("Output folder (auto)", "output", "directory"),
            ("YAML config (optional)", "config", "config"),
            ("MATLAB (Dataset only)", "matlab", "matlab"),
        )
        for row, (label, key, kind) in enumerate(rows):
            ttk.Label(frame, text=label).grid(row=row, column=0, sticky="w", pady=4)
            variable = self.schema_label if key == "schema" else self.fields[key]
            entry = ttk.Entry(frame, textvariable=variable, width=65)
            entry.grid(row=row, column=1, sticky="ew", padx=8)
            if key == "schema":
                entry.state(["readonly"])
            self.controls.append(entry)
            button = ttk.Button(
                frame,
                text="File..." if kind in ("blf", "dbc") else "Browse...",
                command=lambda k=key, t=kind: self.browse(k, t),
            )
            button.grid(row=row, column=2)
            self.controls.append(button)
            if key in ("blf", "dbc_folder"):
                folder = ttk.Button(
                    frame, text="Folder...", command=lambda k=key: self.browse(k, "directory")
                )
                folder.grid(row=row, column=3, padx=4)
                self.controls.append(folder)
            if key == "schema":
                clear = ttk.Button(frame, text="Clear", command=self.clear_schema)
                clear.grid(row=row, column=3)
                self.controls.append(clear)
        advanced = ttk.LabelFrame(frame, text="Output and optional Excel mapping", padding=6)
        advanced.grid(row=7, column=0, columnspan=4, sticky="ew", pady=6)
        advanced.columnconfigure(1, weight=1)
        ttk.Label(advanced, text="Output format").grid(row=0, column=0, sticky="w")
        output_format = ttk.Combobox(
            advanced,
            textvariable=self.fields["output_format"],
            values=("mat", "dataset"),
            state="readonly",
            width=15,
        )
        output_format.grid(row=0, column=1, sticky="w")
        self.controls.append(output_format)
        ttk.Label(advanced, text="mat: no MATLAB | dataset: Signal Editor").grid(
            row=0, column=2, columnspan=2
        )
        sheets = ttk.Button(advanced, text="Select sheets...", command=self.select_sheets)
        sheets.grid(row=1, column=0, sticky="w")
        self.controls.append(sheets)
        ttk.Label(advanced, textvariable=self.sheet_label).grid(
            row=1, column=1, columnspan=3, sticky="w"
        )
        for index, (key, label) in enumerate(
            (
                ("element_column", "Element column"),
                ("can_column", "CAN column"),
                ("interface_column", "Interface column"),
                ("dtype_column", "Datatype column"),
                ("header_row", "Header row (1-based)"),
            )
        ):
            row, column = 2 + index // 2, 2 * (index % 2)
            ttk.Label(advanced, text=label).grid(row=row, column=column, sticky="w", padx=4)
            entry = ttk.Entry(advanced, textvariable=self.fields[key])
            entry.grid(row=row, column=column + 1, sticky="ew", padx=4)
            self.controls.append(entry)
        ttk.Label(
            advanced, text="Blank = auto-detect. Columns accept a header name or 1-based number."
        ).grid(row=5, column=0, columnspan=4, sticky="w")
        ttk.Label(
            frame,
            text=(
                "TXT without Excel: leave Excel blank; element names match DBC signals "
                "(all interfaces are inputs).\n"
                "Empty YAML uses signal_mapping.yaml beside the .pyz when present."
            ),
        ).grid(row=8, column=0, columnspan=4, sticky="w", pady=(4, 8))
        actions = ttk.Frame(frame)
        actions.grid(row=9, column=0, columnspan=4, sticky="ew")
        for label, command in (
            ("Convert", self.start_conversion),
            ("Check environment", self.start_check),
        ):
            button = ttk.Button(actions, text=label, command=command)
            button.pack(side="left", padx=(0, 8))
            self.controls.append(button)
        self.progress = ttk.Progressbar(actions, mode="indeterminate")
        self.progress.pack(side="right", fill="x", expand=True)
        ttk.Label(frame, textvariable=self.status).grid(
            row=10, column=0, columnspan=4, sticky="w", pady=6
        )
        self.output = Text(frame, wrap="word", state="disabled", height=16)
        self.output.grid(row=11, column=0, columnspan=3, sticky="nsew")
        scrollbar = ttk.Scrollbar(frame, command=self.output.yview)
        scrollbar.grid(row=11, column=3, sticky="ns")
        self.output.configure(yscrollcommand=scrollbar.set)
        frame.rowconfigure(11, weight=1)
        root.protocol("WM_DELETE_WINDOW", self.close)
        self.poll_id = root.after(75, self.poll)

    def clear_schema(self) -> None:
        self.schema_files = []
        self.schema_label.set("")

    def select_sheets(self) -> None:
        try:
            from openpyxl import load_workbook

            book = load_workbook(self.fields["excel"].get(), read_only=True)
            try:
                names = book.sheetnames
            finally:
                book.close()
        except Exception as exc:
            messagebox.showerror(
                "Excel workbook",
                f"Choose a readable Excel workbook first.\n{exc}",
                parent=self.root,
            )
            return
        window = Toplevel(self.root)
        window.title("Select Excel sheets")
        lists = {}
        for index, side in enumerate(("input", "output")):
            ttk.Label(window, text=side.capitalize() + " sheets (Ctrl-click for multiple)").grid(
                row=0, column=index, padx=8, pady=4
            )
            box = Listbox(window, selectmode="extended", exportselection=False, width=40)
            box.grid(row=1, column=index, padx=8)
            for position, name in enumerate(names):
                box.insert(END, name)
                if name in self.sheets[side]:
                    box.selection_set(position)
            lists[side] = box

        def save() -> None:
            self.sheets = {
                side: [names[index] for index in box.curselection()] for side, box in lists.items()
            }
            self.sheet_label.set(
                " | ".join(
                    f"{side}: {', '.join(selected)}"
                    for side, selected in self.sheets.items()
                    if selected
                )
                or "Auto-detect sheets"
            )
            window.destroy()

        ttk.Button(window, text="Apply (empty = auto)", command=save).grid(
            row=2, column=0, columnspan=2, pady=8
        )
        window.transient(self.root)
        window.grab_set()

    def browse(self, key: str, kind: str) -> None:
        if kind == "directory":
            value = filedialog.askdirectory(parent=self.root, mustexist=key != "output")
        elif kind == "schema":
            files = filedialog.askopenfilenames(
                parent=self.root, filetypes=[("TXT schema", "*.txt")]
            )
            if files:
                self.schema_files = list(files)
                self.schema_label.set("; ".join(files))
            return
        else:
            filters = {
                "excel": [("Excel workbook", "*.xlsx *.xlsm")],
                "blf": [("CAN BLF log", "*.blf")],
                "dbc": [("DBC definitions", "*.dbc *.DBC")],
                "config": [("YAML mapping", "*.yaml *.yml")],
                "matlab": [("MATLAB executable", "matlab.exe matlab"), ("All files", "*")],
            }
            value = filedialog.askopenfilename(parent=self.root, filetypes=filters[kind])
        if value:
            self.fields[key].set(value)

    def append_log(self, text: str) -> None:
        self.output.configure(state="normal")
        self.output.insert(END, text + "\n")
        self.output.see(END)
        self.output.configure(state="disabled")

    def start_conversion(self) -> None:
        if self.busy:
            return
        args = build_arg_parser().parse_args([])
        for name, variable in self.fields.items():
            setattr(args, name, variable.get().strip() or None)
        args.schema = list(self.schema_files)
        args.input_sheet, args.output_sheet = self.sheets["input"], self.sheets["output"]
        try:
            if args.header_row:
                args.header_row = int(args.header_row)
            paths = resolve_paths(args, _GuiParser())
        except Exception as exc:
            self.status.set("Invalid inputs")
            messagebox.showerror("Invalid inputs", str(exc), parent=self.root)
            return
        self.start_job(lambda emit: convert(paths, log_callback=emit), "Converting...")

    def start_check(self) -> None:
        matlab = self.fields["matlab"].get().strip() or None
        require_matlab = self.fields["output_format"].get() == "dataset"
        self.start_job(
            lambda emit: check_environment(matlab, emit, require_matlab=require_matlab),
            "Checking environment...",
        )

    def start_job(self, job: Callable[[Callable[[str], None]], int], label: str) -> None:
        if self.busy:
            return
        self.busy = True
        self.status.set(label)
        for control in self.controls:
            control.state(["disabled"])
        self.progress.start(12)
        events, debug = self.events, self.debug

        def worker() -> None:
            try:
                code = job(lambda text: events.put(("log", text)))
                events.put(("done", code))
            except Exception as exc:
                events.put(("log", f"ERROR: {exc}"))
                if debug:
                    events.put(("log", traceback.format_exc()))
                events.put(("done", 1))

        threading.Thread(target=worker, name="blf-converter", daemon=False).start()

    def poll(self) -> None:
        # Bound each tick so a large conversion log cannot starve the UI.
        for _ in range(200):
            try:
                kind, value = self.events.get_nowait()
            except queue.Empty:
                break
            if kind == "log":
                self.append_log(str(value))
            else:
                self.busy = False
                self.progress.stop()
                self.progress.configure(value=0)
                for control in self.controls:
                    control.state(["!disabled"])
                self.status.set("Completed" if value == 0 else "Failed - see log")
        self.poll_id = self.root.after(75, self.poll)

    def close(self) -> None:
        if self.busy:
            messagebox.showinfo(
                "Operation running",
                "Wait for the operation to finish before closing.",
                parent=self.root,
            )
            return
        self.root.after_cancel(self.poll_id)
        self.root.destroy()


def launch(args: argparse.Namespace | None = None) -> int:
    root = Tk()
    Application(root, args)
    root.mainloop()
    return 0
